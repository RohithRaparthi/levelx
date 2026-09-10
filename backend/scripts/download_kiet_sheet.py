import requests
import os
import openpyxl
import json

url = 'https://docs.google.com/spreadsheets/d/1p0PDZzI4xm060Z3BrBeGIU-V2CUyjnpPKPtyB0bSrbY/export?format=xlsx'
headers = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'
}

resp = requests.get(url, headers=headers, timeout=20)
out_path = 'backend/scripts/kiet_sheet.xlsx'
with open(out_path, 'wb') as f:
    f.write(resp.content)

print(f"Downloaded KIET sheet to {out_path}, size: {os.path.getsize(out_path)} bytes")

wb = openpyxl.load_workbook(out_path)
print("Sheets in workbook:", wb.sheetnames)

sheet_data = {}
for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    data = []
    headers_list = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers_list = [str(cell) if cell is not None else f"col_{j}" for j, cell in enumerate(row)]
        else:
            if any(row):
                row_dict = {headers_list[j]: (str(cell).strip() if cell is not None else None) for j, cell in enumerate(row) if j < len(headers_list)}
                data.append(row_dict)
    sheet_data[sheetname] = data
    print(f"Sheet '{sheetname}': {len(data)} rows")

with open('backend/scripts/kiet_raw.json', 'w', encoding='utf-8') as f:
    json.dump(sheet_data, f, indent=2)

print("Saved raw sheet data to backend/scripts/kiet_raw.json")
