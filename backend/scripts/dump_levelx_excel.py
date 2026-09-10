import openpyxl
import json

wb = openpyxl.load_workbook('backend/scripts/LevelX.xlsx')

data = {}
for sheet in wb.sheetnames:
    ws = wb[sheet]
    rows = []
    headers = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            headers = [str(c).strip() if c is not None else f"col_{j}" for j, c in enumerate(row)]
        else:
            if any(row):
                row_dict = {headers[j]: (str(c).strip() if c is not None else None) for j, c in enumerate(row) if j < len(headers)}
                rows.append(row_dict)
    data[sheet] = rows

with open('backend/scripts/kiet_excel_dump.json', 'w', encoding='utf-8') as f:
    json.dump(data, f, indent=2)

print("Saved kiet_excel_dump.json")
for k, v in data.items():
    print(f"Sheet '{k}': {len(v)} rows")
